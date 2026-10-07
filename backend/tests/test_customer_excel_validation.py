"""POST /customer-validation/excel: filled Customer Detail Excel vs the GST registry.

No real provider is ever called: `verify_gstin` is replaced in every test, and
any outbound `requests` call fails the test. Workbooks are synthetic (dummy
values) and built in tmp memory with the Customer Detail labels; the real
sample file is never used here.
"""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest

from app.config.config import settings
from app.services import customer_excel_validation as cev
from app.services.gstin_verification import GstinVerificationResult

URL = "/customer-validation/excel"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

GSTIN = "04AAACA0000A1Z5"      # dummy, well-formed
PAN = "AAACA0000A"

# The fixed Customer Detail layout: (row, label), labels in B, values in C.
LABELS = [
    (3, "No."), (4, "Name"), (5, "Search Name"), (6, "PAN Number"), (7, "GST Number"),
    (8, "Address"), (9, "Address 2"), (10, "State"), (11, "City"), (12, "Contact"),
    (13, "Phone No."), (14, "Supply Type code"), (15, "Ship-to Code"),
    (16, "Customer Posting Group"), (17, "Currency Code"), (18, "Payment Terms Code"),
    (19, "Salesperson Code"), (20, "Country/Region Code"), (21, "Location Code"),
    (22, "Gen. Bus. Posting Group"), (23, "ZIP Code"), (24, "State"), (25, "Email"),
    (26, "Email Cc"), (27, "Tax Area Code"), (28, "Tax Liable"),
]

DUMMY = {
    3: "CUST/9999", 4: "Example Media Labs Pvt. Ltd.", 5: "EXAMPLE MEDIA LABS PVT. LTD.",
    6: PAN, 7: GSTIN, 8: "1st Floor, Plot No 12", 9: "Example House, IT Park",
    10: "CHANDIGARH", 11: "CHANDIGARH", 12: "Test Contact", 13: "9000000000", 14: "B2B",
    15: "04", 16: "LEASING", 18: "6 DAYS", 20: "IN", 22: "DOM", 23: 160101, 24: "Chandigarh",
    25: "test@example.invalid", 27: "GST18 (UT)", 28: True,
}

REGISTRY = dict(
    checked=True, active=True, record_found=True, provider="decentro", document_type="GSTIN_DETAILED",
    status="Active", legal_name="EXAMPLE MEDIA LABS PRIVATE LIMITED", trade_name="EXAMPLE MEDIA LABS",
    pan=PAN, address="1st Floor, Plot No 12, Example House, IT Park, Chandigarh, Chandigarh, 160101",
    city="Chandigarh", state="Chandigarh", pincode="160101",
    constitution_of_business="Private Limited Company", taxpayer_type="Regular",
    registration_date="01/07/2017", filing_status=[{"filing_year": "2024-25"}] * 50,
    directors=[{"name": "SOMEONE", "begin_date": "", "end_date": ""}],
)


def workbook_bytes(values: dict | None = None, labels=LABELS, sheet="Indian Company", extra_sheets=()) -> bytes:
    values = DUMMY if values is None else values
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for name in extra_sheets:
        wb.create_sheet(name)
    for row, label in labels:
        ws[f"B{row}"] = label
        if row in values and values[row] is not None:
            ws[f"C{row}"] = values[row]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def with_values(changes: dict) -> dict:
    v = dict(DUMMY)
    v.update(changes)
    return v


@pytest.fixture(autouse=True)
def _no_network(monkeypatch):
    import requests

    def _blocked(*a, **k):
        raise AssertionError("outbound HTTP during a customer-validation test")

    monkeypatch.setattr(requests.sessions.Session, "request", _blocked)
    assert settings.GSTIN_API_ENABLED is False


@pytest.fixture
def registry(monkeypatch):
    """Replace verify_gstin; returns the call log and lets a test set the result."""
    state = {"calls": [], "result": GstinVerificationResult(**REGISTRY)}

    def fake(gstin):
        state["calls"].append(gstin)
        return state["result"]

    monkeypatch.setattr(cev, "verify_gstin", fake)
    return state


def post(client, content: bytes, name="customer.xlsx", ctype=XLSX):
    return client.post(URL, files={"file": (name, content, ctype)})


def rows_by_label(body) -> dict:
    return {r["label"]: r for r in body["validations"]}


def status_of(body, label) -> str:
    return rows_by_label(body)[label]["status"]


# ---------------------------------------------------------------------------
# upload / file handling
# ---------------------------------------------------------------------------

class TestUpload:
    def test_valid_upload_full_shape(self, auth_client, registry):
        r = post(auth_client, workbook_bytes())
        assert r.status_code == 200, r.text
        body = r.json()
        assert set(body) == {"file_name", "layout", "sheet", "extracted_data", "gst_verification",
                             "validations", "summary"}
        assert body["layout"] == "customer_detail_v1" and body["sheet"] == "Indian Company"
        assert body["extracted_data"]["gstin"] == GSTIN
        assert body["extracted_data"]["zip_code"] == "160101"           # numeric cell, no ".0"
        assert len(body["validations"]) == len(LABELS)
        assert sum(body["summary"].values()) == len(LABELS)
        assert registry["calls"] == [GSTIN]                               # exactly one GST call

    def test_registry_block_is_normalized_not_raw(self, auth_client, registry):
        gst = post(auth_client, workbook_bytes()).json()["gst_verification"]
        assert (gst["provider"], gst["document_type"], gst["status"]) == ("decentro", "GSTIN_DETAILED", "Active")
        reg = gst["registry"]
        assert reg["legal_name"] == REGISTRY["legal_name"]
        for left_out in ("filing_status", "directors", "business_details", "reference_id",
                         "provider_reference_id", "kycResult"):
            assert left_out not in reg and left_out not in gst

    @pytest.mark.parametrize("name,content,code", [
        ("customer.csv", b"a,b,c", 415),
        ("customer.xls", b"\xd0\xcf\x11\xe0legacy", 415),
        ("customer.xlsx", b"this is not a zip", 422),
        ("customer.xlsx", b"PK\x03\x04 corrupt zip body", 422),
    ])
    def test_non_excel_or_corrupt_rejected(self, auth_client, registry, name, content, code):
        r = post(auth_client, content, name=name)
        assert r.status_code == code
        assert registry["calls"] == []

    def test_oversize_rejected(self, auth_client, registry, monkeypatch):
        monkeypatch.setattr(cev, "MAX_UPLOAD_BYTES", 1000)
        import app.routers.customer_validation as router_mod
        monkeypatch.setattr(router_mod, "MAX_UPLOAD_BYTES", 1000)
        r = post(auth_client, workbook_bytes())
        assert r.status_code == 413
        assert registry["calls"] == []

    def test_xlsx_that_is_a_different_layout_is_rejected(self, auth_client, registry):
        moved = [(row + 1, label) for row, label in LABELS]          # everything one row down
        r = post(auth_client, workbook_bytes(labels=moved))
        assert r.status_code == 422
        assert "does not match the Customer Detail layout" in r.json()["detail"]
        assert registry["calls"] == []

    def test_renamed_label_is_rejected(self, auth_client, registry):
        labels = [(row, "GSTIN" if label == "GST Number" else label) for row, label in LABELS]
        r = post(auth_client, workbook_bytes(labels=labels))
        assert r.status_code == 422 and "B7" in r.json()["detail"]

    def test_matching_sheet_found_even_if_not_first(self, auth_client, registry):
        wb = openpyxl.Workbook()
        wb.active.title = "Notes"
        ws = wb.create_sheet("Some Other Name")
        for row, label in LABELS:
            ws[f"B{row}"] = label
            if row in DUMMY:
                ws[f"C{row}"] = DUMMY[row]
        buf = io.BytesIO(); wb.save(buf)
        r = post(auth_client, buf.getvalue())
        assert r.status_code == 200 and r.json()["sheet"] == "Some Other Name"

    def test_labels_tolerate_case_spacing_and_colon(self, auth_client, registry):
        labels = [(row, f"  {label.upper()} :") for row, label in LABELS]
        assert post(auth_client, workbook_bytes(labels=labels)).status_code == 200

    def test_temp_file_is_always_deleted(self, auth_client, registry, monkeypatch, tmp_path):
        import tempfile
        monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
        post(auth_client, workbook_bytes())
        post(auth_client, b"PK\x03\x04 corrupt", name="x.xlsx")
        assert list(Path(tmp_path).glob("customer-validation-*")) == []


# ---------------------------------------------------------------------------
# GSTIN handling
# ---------------------------------------------------------------------------

class TestGstin:
    def test_missing_gstin_no_call_everything_not_validated(self, auth_client, registry):
        body = post(auth_client, workbook_bytes(with_values({7: None}))).json()
        assert registry["calls"] == []
        assert body["gst_verification"]["attempted"] is False
        assert body["gst_verification"]["error"] == "GSTIN missing in Excel"
        assert status_of(body, "GST Number") == "NOT_PRESENT_IN_EXCEL"
        for label in ("Name", "PAN Number", "Address", "State (row 10)", "City", "ZIP Code"):
            assert status_of(body, label) == "NOT_VALIDATED", label

    def test_malformed_gstin_no_call(self, auth_client, registry):
        body = post(auth_client, workbook_bytes(with_values({7: "04ABC123"}))).json()
        assert registry["calls"] == []
        assert "not a valid 15-character GSTIN" in body["gst_verification"]["error"]
        assert status_of(body, "GST Number") == "NOT_VALIDATED"
        assert status_of(body, "Name") == "NOT_VALIDATED"

    def test_gstin_lowercase_in_excel_is_normalized(self, auth_client, registry):
        post(auth_client, workbook_bytes(with_values({7: GSTIN.lower()})))
        assert registry["calls"] == [GSTIN]

    def test_active_gstin_matches(self, auth_client, registry):
        row = rows_by_label(post(auth_client, workbook_bytes()).json())["GST Number"]
        assert row["status"] == "MATCH" and "Active" in row["detail"]

    def test_cancelled_gstin_needs_review(self, auth_client, registry):
        registry["result"] = GstinVerificationResult(**{**REGISTRY, "active": False, "status": "Cancelled"})
        body = post(auth_client, workbook_bytes()).json()
        row = rows_by_label(body)["GST Number"]
        assert row["status"] == "REVIEW" and "Cancelled" in row["detail"]
        assert body["gst_verification"]["status"] == "Cancelled"
        assert status_of(body, "Name") == "MATCH"      # the registry record is still compared

    def test_gstin_not_found_is_mismatch_rest_not_validated(self, auth_client, registry):
        registry["result"] = GstinVerificationResult(
            checked=True, record_found=False, active=False, status="Not found", provider="decentro",
            document_type="GSTIN_DETAILED")
        body = post(auth_client, workbook_bytes()).json()
        assert status_of(body, "GST Number") == "MISMATCH"
        assert status_of(body, "Name") == "NOT_VALIDATED"
        assert body["gst_verification"]["registry"] == {}

    def test_both_providers_down(self, auth_client, registry):
        registry["result"] = GstinVerificationResult(
            checked=False, error="Decentro server_error (HTTP 503); fallback: gstinapi.in unexpected HTTP 500",
            primary_error="Decentro server_error (HTTP 503)")
        body = post(auth_client, workbook_bytes()).json()
        assert status_of(body, "GST Number") == "NOT_VALIDATED"
        compared = ("Name", "Search Name", "PAN Number", "Address", "Address 2",
                    "State (row 10)", "City", "ZIP Code", "State (row 24)")
        assert all(status_of(body, label) == "NOT_VALIDATED" for label in compared)
        assert body["gst_verification"]["checked"] is False
        assert "Decentro server_error" in body["gst_verification"]["error"]


# ---------------------------------------------------------------------------
# field comparisons
# ---------------------------------------------------------------------------

class TestComparisons:
    def test_happy_path_identity_fields_match(self, auth_client, registry):
        body = post(auth_client, workbook_bytes()).json()
        for label in ("Name", "Search Name", "PAN Number", "GST Number", "Address", "Address 2",
                      "State (row 10)", "State (row 24)", "City", "ZIP Code"):
            assert status_of(body, label) == "MATCH", (label, rows_by_label(body)[label])

    def test_name_suffix_normalization_matches_legal_name(self, auth_client, registry):
        row = rows_by_label(post(auth_client, workbook_bytes()).json())["Name"]
        assert row["status"] == "MATCH" and "legal name" in row["detail"]

    def test_name_equal_to_trade_name_matches(self, auth_client, registry):
        row = rows_by_label(post(auth_client, workbook_bytes(with_values({4: "Example Media Labs"}))).json())["Name"]
        assert row["status"] == "MATCH" and "trade name" in row["detail"]

    def test_name_typo_is_review(self, auth_client, registry):
        row = rows_by_label(post(auth_client, workbook_bytes(with_values({4: "Exampel Media Labs Pvt Ltd"}))).json())["Name"]
        assert row["status"] == "REVIEW"

    def test_different_name_is_mismatch(self, auth_client, registry):
        row = rows_by_label(post(auth_client, workbook_bytes(with_values({4: "Totally Different Traders"}))).json())["Name"]
        assert row["status"] == "MISMATCH"

    def test_pan_mismatch(self, auth_client, registry):
        assert status_of(post(auth_client, workbook_bytes(with_values({6: "ZZZZZ9999Z"}))).json(), "PAN Number") == "MISMATCH"

    def test_state_mismatch_and_alias_match(self, auth_client, registry):
        body = post(auth_client, workbook_bytes(with_values({10: "Punjab", 24: "chandigarh"}))).json()
        assert status_of(body, "State (row 10)") == "MISMATCH"
        assert status_of(body, "State (row 24)") == "MATCH"

    def test_pin_mismatch(self, auth_client, registry):
        assert status_of(post(auth_client, workbook_bytes(with_values({23: "160055"}))).json(), "ZIP Code") == "MISMATCH"

    def test_address_low_overlap_is_review_never_mismatch(self, auth_client, registry):
        body = post(auth_client, workbook_bytes(with_values({8: "Shop 4, Some Other Market", 9: "Elsewhere Road"}))).json()
        assert status_of(body, "Address") == "REVIEW"
        assert status_of(body, "Address 2") == "REVIEW"
        assert "MISMATCH" not in {status_of(body, "Address"), status_of(body, "Address 2")}

    def test_city_not_in_address_is_review(self, auth_client, registry):
        assert status_of(post(auth_client, workbook_bytes(with_values({11: "Mohali"}))).json(), "City") == "REVIEW"

    def test_empty_excel_cell_with_registry_value(self, auth_client, registry):
        body = post(auth_client, workbook_bytes(with_values({6: None, 23: None}))).json()
        assert status_of(body, "PAN Number") == "NOT_PRESENT_IN_EXCEL"
        assert status_of(body, "ZIP Code") == "NOT_PRESENT_IN_EXCEL"

    def test_non_registry_fields_are_never_mismatch(self, auth_client, registry):
        body = post(auth_client, workbook_bytes()).json()
        for label in ("Contact", "Phone No.", "Email", "Payment Terms Code", "Country/Region Code",
                      "Tax Area Code", "Tax Liable", "Customer Posting Group"):
            assert status_of(body, label) == "NOT_AVAILABLE_FROM_GST_API", label
        assert status_of(body, "No.") == "NOT_VALIDATED"

    def test_fallback_provider_missing_fields_are_not_available(self, auth_client, registry):
        registry["result"] = GstinVerificationResult(
            checked=True, active=True, record_found=True, provider="gstinapi", status="Active",
            legal_name=REGISTRY["legal_name"], trade_name="", address="1 Some Road",
            city="Chandigarh", pincode="160101", primary_error="Decentro auth_rejected (HTTP 401)")
        body = post(auth_client, workbook_bytes()).json()
        assert status_of(body, "PAN Number") == "NOT_AVAILABLE_FROM_GST_API"
        assert status_of(body, "State (row 10)") == "NOT_AVAILABLE_FROM_GST_API"
        assert status_of(body, "Name") == "MATCH"
        assert body["gst_verification"]["provider"] == "gstinapi"
        assert body["gst_verification"]["primary_error"] == "Decentro auth_rejected (HTTP 401)"


# ---------------------------------------------------------------------------
# auth, isolation, side effects
# ---------------------------------------------------------------------------

class TestAuthAndSideEffects:
    def test_requires_authentication(self, client, registry):
        r = client.post(URL, files={"file": ("c.xlsx", workbook_bytes(), XLSX)})
        assert r.status_code == 401
        assert registry["calls"] == []

    def test_creates_no_customer_or_vendor_and_no_bc_payload(self, auth_client, registry, monkeypatch):
        import app.services.bc_mapper as bc_mapper

        def _never(*a, **k):
            raise AssertionError("BC payload must not be built by the validation flow")

        monkeypatch.setattr(bc_mapper, "customer_to_bc_payload", _never)
        before = (auth_client.get("/customers").json(), auth_client.get("/vendors").json())
        assert post(auth_client, workbook_bytes()).status_code == 200
        after = (auth_client.get("/customers").json(), auth_client.get("/vendors").json())
        assert before == after == ([], [])

    def test_response_never_contains_provider_secrets(self, auth_client, registry, monkeypatch):
        monkeypatch.setattr(settings, "DECENTRO_CLIENT_SECRET", "SECRET-client-VALUE-123")
        monkeypatch.setattr(settings, "GSTIN_API_KEY", "SECRET-apikey-VALUE-789")
        text = post(auth_client, workbook_bytes()).text
        assert "SECRET-client-VALUE-123" not in text and "SECRET-apikey-VALUE-789" not in text

    def test_layout_yaml_is_isolated_from_vendor_mappings(self):
        from app.services.extraction_pipeline.excel.excel_mapper import ExcelMapper
        assert "customer_detail_v1" not in ExcelMapper.available()

    def test_layout_labels_match_the_confirmed_form(self):
        _, fields, sheet = cev.load_layout()
        assert sheet == "Indian Company"
        assert [(f.row, f.label) for f in fields] == LABELS
        assert all(f.cell == f"C{f.row}" and f.label_cell == f"B{f.row}" for f in fields)
