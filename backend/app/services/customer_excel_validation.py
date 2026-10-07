"""Validate a filled Customer Detail Excel against the live GST registry.

A separate, read-only flow (POST /customer-validation/excel):

    filled Excel -> layout check -> read values -> verify_gstin() (Decentro
    GSTIN_DETAILED, gstinapi.in fallback -- the EXISTING service, called at most
    once) -> compare Excel values with the registry -> structured result

Nothing is persisted: the upload lives in a temp file deleted before
returning, no customer/vendor record is created, nothing is pushed to Business
Central. The existing document -> review -> POST /customers flow is untouched.

The layout comes from config/excel_validation/customer_detail_v1.yaml. Every
column-B label is checked before any value is trusted; a workbook whose labels
don't match is rejected, never guessed at.

Statuses per field:
    MATCH                       Excel agrees with the registry
    REVIEW                      close but not equal (fuzzy name, partial address
                                overlap, city not literally in the address,
                                GSTIN registered but not active)
    MISMATCH                    registry clearly disagrees
    NOT_AVAILABLE_FROM_GST_API  the registry has no such field (or the fallback
                                provider didn't return it) -- never a mismatch
    NOT_PRESENT_IN_EXCEL        the registry has it, the Excel cell is empty
    NOT_VALIDATED               the GST check didn't run or didn't answer
"""

from __future__ import annotations

import io
import logging
import os
import re
import tempfile
import zipfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import openpyxl
import yaml
from openpyxl.utils.cell import coordinate_from_string, get_column_letter, column_index_from_string
from rapidfuzz import fuzz

from app.config.config import BACKEND_DIR
from app.services.extraction_pipeline.excel.excel_mapper import ExcelMapper
from app.services.extraction_pipeline.extract.address_lookups import canonical_city, canonical_state
from app.services.gstin_verification import GstinVerificationResult, verify_gstin

logger = logging.getLogger(__name__)

LAYOUT_DIR = BACKEND_DIR / "config" / "excel_validation"
LAYOUT_NAME = "customer_detail_v1"

ALLOWED_EXTENSIONS = (".xlsx", ".xlsm")
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
_ZIP_MAGIC = b"PK\x03\x04"           # .xlsx / .xlsm are zip packages

MATCH = "MATCH"
REVIEW = "REVIEW"
MISMATCH = "MISMATCH"
NOT_AVAILABLE = "NOT_AVAILABLE_FROM_GST_API"
NOT_PRESENT = "NOT_PRESENT_IN_EXCEL"
NOT_VALIDATED = "NOT_VALIDATED"
STATUSES = (MATCH, REVIEW, MISMATCH, NOT_AVAILABLE, NOT_PRESENT, NOT_VALIDATED)

NAME_FUZZY_THRESHOLD = 85            # same threshold validator.py uses across documents
ADDRESS_OVERLAP_MATCH = 0.6

_GSTIN_RE = re.compile(r"^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]$")

# Fields compared with the registry. Everything else in the layout is shown
# but has no GST registry counterpart.
_NAME_FIELDS = ("name", "search_name")
_ADDRESS_LINE_FIELDS = ("address_1", "address_2")
_STATE_FIELDS = ("state", "state_2")
_COMPARED_FIELDS = {"gstin", "pan", "city", "zip_code", *_NAME_FIELDS, *_ADDRESS_LINE_FIELDS, *_STATE_FIELDS}
_NOT_A_REGISTRY_FIELD = {"customer_no"}   # an internal BC number: shown, never validated


class ExcelValidationError(ValueError):
    """The upload can't be validated (wrong type, too large, unreadable,
    unknown layout). The router turns it into an HTTP error."""

    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.status_code = status_code


# --- layout -----------------------------------------------------------------

@dataclass(frozen=True)
class LayoutField:
    field: str
    cell: str
    label: str
    label_cell: str      # the column-B cell holding the label for this value
    row: int


def load_layout(name: str = LAYOUT_NAME) -> tuple[ExcelMapper, list[LayoutField], Optional[str]]:
    """The ExcelMapper for reading values, the per-field labels, and the
    preferred sheet name. Labels sit one column left of each value cell."""
    path = LAYOUT_DIR / f"{name}.yaml"
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    mapper = ExcelMapper.load(name, directory=LAYOUT_DIR)
    fields: list[LayoutField] = []
    for field, cfg in (data.get("fields") or {}).items():
        cell = str(cfg["cell"]).upper()
        col, row = coordinate_from_string(cell)
        label_col = get_column_letter(column_index_from_string(col) - 1)
        fields.append(LayoutField(field, cell, str(cfg["label"]), f"{label_col}{row}", row))
    return mapper, fields, data.get("sheet")


def _norm_label(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().rstrip(":").strip().casefold()


def _label_mismatches(sheet, fields: list[LayoutField]) -> list[str]:
    return [
        f"{f.label_cell} should be {f.label!r}"
        for f in fields
        if _norm_label(sheet[f.label_cell].value) != _norm_label(f.label)
    ]


def _pick_sheet(path: Path, fields: list[LayoutField], preferred: Optional[str]) -> str:
    """The sheet whose labels match the layout: the preferred one if it
    matches, otherwise the first that does. Raises if none does."""
    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:  # corrupt package, not really a workbook, ...
        raise ExcelValidationError(
            f"The file could not be opened as an Excel workbook ({type(exc).__name__})."
        ) from exc
    try:
        names = workbook.sheetnames
        order = ([preferred] if preferred in names else []) + [n for n in names if n != preferred]
        first_problems: list[str] = []
        for name in order:
            problems = _label_mismatches(workbook[name], fields)
            if not problems:
                return name
            if not first_problems:
                first_problems = problems
        raise ExcelValidationError(
            "This Excel does not match the Customer Detail layout "
            f"(labels in column B, values in column C, rows {fields[0].row}-{fields[-1].row}). "
            f"First differences: {'; '.join(first_problems[:3])}."
        )
    finally:
        workbook.close()


# --- comparison helpers -------------------------------------------------------

_NAME_EQUIVALENTS = (
    (r"\bPRIVATE\b", "PVT"),
    (r"\bLIMITED\b", "LTD"),
    (r"\bCOMPANY\b", "CO"),
    (r"&", " AND "),
)


def normalize_name(value: str) -> str:
    s = (value or "").upper()
    for pattern, repl in _NAME_EQUIVALENTS:
        s = re.sub(pattern, repl, s)
    s = re.sub(r"[^A-Z0-9 ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _tokens(text: str) -> list[str]:
    return [t for t in re.split(r"[^a-z0-9]+", (text or "").casefold()) if t]


def _norm_text(text: str) -> str:
    return " ".join(_tokens(text))


def _row(field: LayoutField, excel: str, status: str, verified: str = "", detail: str = "") -> dict:
    return {
        "field": field.field,
        "label": field.label,
        "row": field.row,
        "excel_value": excel,
        "verified_value": verified,
        "status": status,
        "detail": detail,
    }


def _compare_name(excel: str, legal: str, trade: str) -> tuple[str, str, str]:
    candidates = [(kind, value) for kind, value in (("legal name", legal), ("trade name", trade)) if value]
    if not candidates:
        return NOT_AVAILABLE, "", "registry returned no legal or trade name"
    target = normalize_name(excel)
    for kind, value in candidates:
        if normalize_name(value) == target:
            return MATCH, value, f"equals registry {kind}"
    scored = sorted(((fuzz.ratio(target, normalize_name(v)), kind, v) for kind, v in candidates), reverse=True)
    score, kind, value = scored[0]
    if score >= NAME_FUZZY_THRESHOLD:
        return REVIEW, value, f"close to registry {kind} (similarity {score:.0f}%), not identical"
    return MISMATCH, value, f"differs from registry {kind} (similarity {score:.0f}%)"


def _compare_address_line(excel: str, registry_address: str) -> tuple[str, str]:
    words = _tokens(excel)
    if not words:
        return REVIEW, "no comparable words"
    registry_words = set(_tokens(registry_address))
    found = sum(1 for w in words if w in registry_words)
    share = found / len(words)
    status = MATCH if share >= ADDRESS_OVERLAP_MATCH else REVIEW
    return status, f"{found} of {len(words)} words ({share:.0%}) found in the registry principal address"


def _canonical_state(value: str) -> str:
    return (canonical_state(value) or value or "").strip().casefold()


def _compare_city(excel: str, registry_city: str, registry_address: str) -> tuple[str, str]:
    city = _norm_text(excel)
    if registry_city and city == _norm_text(registry_city):
        return MATCH, "equals registry city"
    canon = canonical_city(excel)
    if canon and registry_city and _norm_text(canon) == _norm_text(registry_city):
        return MATCH, "same city as registry (alias)"
    if city and f" {city} " in f" {_norm_text(registry_address)} ":
        return MATCH, "found in the registry principal address"
    return REVIEW, "not found literally in the registry address (may be a district or alias)"


def _digits(value: str) -> str:
    return re.sub(r"\D", "", value or "")


# --- the comparison ---------------------------------------------------------------

def compare_with_registry(
    values: dict[str, str],
    fields: list[LayoutField],
    verification: Optional[GstinVerificationResult],
    gstin_problem: str = "",
) -> list[dict]:
    """One row per layout field, in layout order. `verification` is None when
    the GST check was not attempted (missing / malformed GSTIN)."""
    registry_ok = bool(verification and verification.checked and verification.record_found)
    duplicate_labels = {label for label, n in Counter(f.label for f in fields).items() if n > 1}
    rows: list[dict] = []

    for f in fields:
        excel = (values.get(f.field) or "").strip()
        display = f if f.label not in duplicate_labels else LayoutField(
            f.field, f.cell, f"{f.label} (row {f.row})", f.label_cell, f.row)

        if f.field in _NOT_A_REGISTRY_FIELD:
            rows.append(_row(display, excel, NOT_VALIDATED, detail="internal reference; not a GST registry field"))
            continue
        if f.field not in _COMPARED_FIELDS:
            rows.append(_row(display, excel, NOT_AVAILABLE, detail="the GST registry has no such field"))
            continue

        if f.field == "gstin":
            rows.append(_gstin_row(display, excel, verification, gstin_problem))
            continue

        if not registry_ok:
            reason = gstin_problem or (verification.error if verification and not verification.checked
                                       else "GSTIN not found in the GST registry")
            rows.append(_row(display, excel, NOT_VALIDATED, detail=f"GST check did not run: {reason}"))
            continue

        v = verification
        if f.field in _NAME_FIELDS:
            if not excel:
                rows.append(_row(display, excel, NOT_PRESENT, v.legal_name or v.trade_name, "empty in Excel"))
            else:
                status, verified, detail = _compare_name(excel, v.legal_name, v.trade_name)
                rows.append(_row(display, excel, status, verified, detail))
        elif f.field == "pan":
            rows.append(_pan_row(display, excel, v))
        elif f.field in _ADDRESS_LINE_FIELDS:
            if not v.address:
                rows.append(_row(display, excel, NOT_AVAILABLE, detail="registry returned no address"))
            elif not excel:
                rows.append(_row(display, excel, NOT_PRESENT, v.address, "empty in Excel"))
            else:
                status, detail = _compare_address_line(excel, v.address)
                rows.append(_row(display, excel, status, v.address, detail))
        elif f.field in _STATE_FIELDS:
            if not v.state:
                rows.append(_row(display, excel, NOT_AVAILABLE, detail="registry returned no state"))
            elif not excel:
                rows.append(_row(display, excel, NOT_PRESENT, v.state, "empty in Excel"))
            elif _canonical_state(excel) == _canonical_state(v.state):
                rows.append(_row(display, excel, MATCH, v.state, "same state as registry"))
            else:
                rows.append(_row(display, excel, MISMATCH, v.state, "different state from registry"))
        elif f.field == "city":
            if not (v.city or v.address):
                rows.append(_row(display, excel, NOT_AVAILABLE, detail="registry returned no city or address"))
            elif not excel:
                rows.append(_row(display, excel, NOT_PRESENT, v.city, "empty in Excel"))
            else:
                status, detail = _compare_city(excel, v.city, v.address)
                rows.append(_row(display, excel, status, v.city or v.address, detail))
        elif f.field == "zip_code":
            if not v.pincode:
                rows.append(_row(display, excel, NOT_AVAILABLE, detail="registry returned no PIN code"))
            elif not excel:
                rows.append(_row(display, excel, NOT_PRESENT, v.pincode, "empty in Excel"))
            elif _digits(excel) == _digits(v.pincode):
                rows.append(_row(display, excel, MATCH, v.pincode, "same PIN as registry"))
            else:
                rows.append(_row(display, excel, MISMATCH, v.pincode, "different PIN from registry"))
    return rows


def _gstin_row(f: LayoutField, excel: str, v: Optional[GstinVerificationResult], problem: str) -> dict:
    if not excel:
        return _row(f, excel, NOT_PRESENT, detail="GSTIN missing in Excel; GST check not run")
    if problem:
        return _row(f, excel, NOT_VALIDATED, detail=problem)
    if v is None or not v.checked:
        return _row(f, excel, NOT_VALIDATED,
                    detail=f"GST check failed: {v.error if v else 'not attempted'}")
    if not v.record_found:
        return _row(f, excel, MISMATCH, detail="GSTIN not found in the GST registry")
    if not v.active:
        return _row(f, excel, REVIEW, excel,
                    f"registered, but GST status is {v.status or 'not Active'}")
    return _row(f, excel, MATCH, excel, f"registered; GST status {v.status or 'Active'}")


def _pan_row(f: LayoutField, excel: str, v: GstinVerificationResult) -> dict:
    if not v.pan:
        return _row(f, excel, NOT_AVAILABLE, detail="registry returned no PAN")
    if not excel:
        return _row(f, excel, NOT_PRESENT, v.pan, "empty in Excel")
    if excel.strip().upper() == v.pan.strip().upper():
        return _row(f, excel, MATCH, v.pan, "same PAN as registry")
    return _row(f, excel, MISMATCH, v.pan, "different PAN from registry")


# --- the response shape ---------------------------------------------------------

_REGISTRY_FIELDS = (
    "legal_name", "trade_name", "pan", "address", "city", "state", "pincode",
    "constitution_of_business", "taxpayer_type", "registration_date", "registration_type",
    "nature_of_business", "nature_of_core_business_activity", "state_jurisdiction",
    "central_jurisdiction", "annual_aggregate_turnover", "mandatory_e_invoicing",
)


def _gst_block(v: Optional[GstinVerificationResult], problem: str) -> dict:
    """Normalized, application-level view of the verification. Never the raw
    provider response; filing history / directors / business details are left
    out on purpose."""
    if v is None:
        return {"attempted": False, "checked": False, "active": False, "record_found": False,
                "provider": "", "document_type": "", "status": "", "primary_error": "",
                "error": problem, "registry": {}}
    data = v.to_dict()
    return {
        "attempted": True,
        "checked": v.checked,
        "active": v.active,
        "record_found": v.record_found,
        "provider": v.provider,
        "document_type": v.document_type,
        "status": v.status,
        "primary_error": v.primary_error,
        "error": v.error,
        "registry": {k: data[k] for k in _REGISTRY_FIELDS} if v.checked and v.record_found else {},
    }


def _gstin_problem(gstin: str) -> str:
    if not gstin:
        return "GSTIN missing in Excel"
    if not _GSTIN_RE.match(gstin):
        return "GSTIN in Excel is not a valid 15-character GSTIN format"
    return ""


# --- entry point --------------------------------------------------------------------

def check_upload(filename: str, content: bytes) -> None:
    """Cheap checks before anything touches openpyxl."""
    if not (filename or "").lower().endswith(ALLOWED_EXTENSIONS):
        raise ExcelValidationError("Only .xlsx or .xlsm Excel files are accepted.", 415)
    if len(content) > MAX_UPLOAD_BYTES:
        raise ExcelValidationError(
            f"The file is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.", 413)
    if not content.startswith(_ZIP_MAGIC) or not zipfile.is_zipfile(io.BytesIO(content)):
        raise ExcelValidationError("The file is not a valid Excel workbook.")


def validate_customer_excel(filename: str, content: bytes) -> dict:
    """Validate one uploaded Customer Detail workbook. Read-only and stateless:
    the bytes go to a temp file that is always deleted; at most ONE GST
    verification call is made, none when the GSTIN is missing or malformed."""
    check_upload(filename, content)
    mapper, fields, preferred_sheet = load_layout()

    suffix = Path(filename).suffix.lower()
    fd, tmp_name = tempfile.mkstemp(prefix="customer-validation-", suffix=suffix)
    tmp_path = Path(tmp_name)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(content)
        sheet = _pick_sheet(tmp_path, fields, preferred_sheet)
        try:
            values = mapper.read_values(str(tmp_path), sheet_name=sheet)
        except Exception as exc:
            raise ExcelValidationError(
                f"The workbook could not be read ({type(exc).__name__}).") from exc
    finally:
        tmp_path.unlink(missing_ok=True)

    extracted = {f.field: values.get(f.field, "") for f in fields}
    gstin = (extracted.get("gstin") or "").strip().upper()
    problem = _gstin_problem(gstin)
    verification = verify_gstin(gstin) if not problem else None

    validations = compare_with_registry(extracted, fields, verification, problem)
    summary = {status: 0 for status in STATUSES}
    for row in validations:
        summary[row["status"]] += 1

    logger.info(
        "customer excel validated: layout=%s sheet=%s gst_attempted=%s provider=%s summary=%s",
        LAYOUT_NAME, sheet, verification is not None,
        verification.provider if verification else "-",
        {k: v for k, v in summary.items() if v},
    )
    return {
        "file_name": Path(filename).name,
        "layout": LAYOUT_NAME,
        "sheet": sheet,
        "extracted_data": extracted,
        "gst_verification": _gst_block(verification, problem),
        "validations": validations,
        "summary": summary,
    }
